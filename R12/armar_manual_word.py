"""
Generador local de manuales en Word (.docx) a partir de un CSV e imágenes.

Requisitos:
- Python 3.10 o superior
- Instalación: pip install pandas python-docx

Ejemplo de uso:
python armar_manual_word.py --csv plantilla.csv --imagenes ./capturas --titulo "Manual de Facturacion" --salida Manual_Borrador.docx
"""

import argparse
import sys
from pathlib import Path
import pandas as pd
from docx import Document
from docx.shared import Inches

COLUMNAS_REQUERIDAS = ["PASO", "ACCIÓN DETECTADA", "CAPTURA", "PRIVACIDAD"]


def construir_manual(ruta_csv: Path, carpeta_img: Path, titulo: str, ruta_salida: Path) -> None:
    if not ruta_csv.exists():
        print(f"Error: No se encontró el archivo CSV en {ruta_csv}")
        sys.exit(1)

    if ruta_salida.exists():
        print(f"Alto: El archivo de salida '{ruta_salida}' ya existe. Cambia el nombre para no sobrescribirlo.")
        sys.exit(1)

    try:
        df = pd.read_csv(ruta_csv, dtype=str).fillna("")
    except Exception as exc:
        print(f"Error al leer el archivo CSV: {exc}")
        sys.exit(1)

    faltantes = [col for col in COLUMNAS_REQUERIDAS if col not in df.columns]
    if faltantes:
        print(f"Error: Faltan las siguientes columnas en el CSV: {', '.join(faltantes)}")
        sys.exit(1)

    doc = Document()
    doc.add_heading(titulo, level=0)
    doc.add_paragraph(
        "BORRADOR PARA REVISIÓN INTERNA. Verifique que todas las capturas "
        "tengan ocultos los datos confidenciales antes de distribuir o exportar a PDF."
    )

    tabla_resumen = doc.add_table(rows=1, cols=4)
    tabla_resumen.style = "Table Grid"
    encabezados = tabla_resumen.rows[0].cells
    for idx, nombre_col in enumerate(COLUMNAS_REQUERIDAS):
        encabezados[idx].text = nombre_col

    for _, fila in df.iterrows():
        celdas = tabla_resumen.add_row().cells
        celdas[0].text = str(fila["PASO"]).strip()
        celdas[1].text = str(fila["ACCIÓN DETECTADA"]).strip()
        celdas[2].text = str(fila["CAPTURA"]).strip()
        celdas[3].text = str(fila["PRIVACIDAD"]).strip()

    doc.add_page_break()
    doc.add_heading("Detalle Paso a Paso", level=1)

    for _, fila in df.iterrows():
        paso = str(fila["PASO"]).strip()
        accion = str(fila["ACCIÓN DETECTADA"]).strip()
        captura = str(fila["CAPTURA"]).strip()
        privacidad = str(fila["PRIVACIDAD"]).strip()

        doc.add_heading(f"{paso}: {accion}", level=2)
        doc.add_paragraph(f"Estado de privacidad reportado: {privacidad}")

        if captura:
            ruta_img = carpeta_img / captura
            if ruta_img.exists() and ruta_img.is_file():
                try:
                    doc.add_picture(str(ruta_img), width=Inches(5.5))
                except Exception as exc:
                    doc.add_paragraph(f"[No se pudo insertar la imagen {captura}: {exc}]")
            else:
                doc.add_paragraph(f"[Imagen pendiente de agregar: {captura}]")

    doc.save(str(ruta_salida))
    print(f"Listo: Manual generado correctamente en '{ruta_salida}'.")


def main() -> None:
    parser = argparse.ArgumentParser(
        description="Crea un manual en Word (.docx) desde una tabla CSV y una carpeta de capturas."
    )
    parser.add_argument("--csv", required=True, help="Ruta al archivo CSV con los pasos.")
    parser.add_argument("--imagenes", default=".", help="Carpeta donde están las capturas de pantalla.")
    parser.add_argument("--titulo", default="Manual de Procedimiento", help="Título principal del manual.")
    parser.add_argument("--salida", required=True, help="Ruta del archivo .docx nuevo a crear.")

    args = parser.parse_args()
    construir_manual(
        ruta_csv=Path(args.csv),
        carpeta_img=Path(args.imagenes),
        titulo=args.titulo,
        ruta_salida=Path(args.salida),
    )


if __name__ == "__main__":
    main()
